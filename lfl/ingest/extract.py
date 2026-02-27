"""
Format-specific content extractors.

Each extractor returns an ExtractedDoc with text, tables, and metadata.
Structured error codes are attached via the `structured_errors` field.
"""

from __future__ import annotations

from pathlib import Path

from .errors import (
    E100_no_content_extracted,
    E101_pdf_stream_corruption,
    E102_ocr_unavailable,
    E103_partial_extraction,
    E104_pymupdf_not_installed,
    E105_pdfplumber_not_installed,
    E106_extractor_exception,
    E110_docx_not_installed,
    E111_pandas_not_installed,
    E112_bs4_not_installed,
    E113_encoding_fallback,
    IngestionError,
)
from .types import ExtractedDoc, ExtractedTable


def extract_text(path: Path, file_type: str) -> ExtractedDoc:
    """
    Extract content from any supported file type.
    
    Dispatches to format-specific extractor.
    """
    extractors = {
        'pdf': extract_pdf,
        'docx': extract_docx,
        'xlsx': extract_xlsx,
        'html': extract_html,
        'text': extract_plain_text,
        'markdown': extract_plain_text,
    }
    
    extractor = extractors.get(file_type, extract_plain_text)
    return extractor(path)


def extract_pdf(path: Path) -> ExtractedDoc:
    """
    Extract text from PDF with fallback chain.
    
    1. Try PyMuPDF (fast)
    2. Fallback to pdfplumber (better tables)
    3. Fallback to OCR (if available and text empty)
    
    Emits structured error codes (LFL-E1xx) for each failure mode.
    """
    text = ""
    tables: list[ExtractedTable] = []
    method = "native"
    warnings: list[str] = []
    errors: list[IngestionError] = []
    page_count = 0
    zlib_errors_detected = False
    
    # Try PyMuPDF first
    try:
        import fitz  # PyMuPDF
        
        doc = fitz.open(path)
        page_count = doc.page_count
        pages = []
        for page in doc:
            pages.append(page.get_text())
        text = "\n\n".join(pages)
        doc.close()
        
        # Detect partial extraction (some pages blank)
        non_empty = sum(1 for p in pages if p.strip())
        if non_empty > 0 and non_empty < page_count:
            err = E103_partial_extraction(str(path), non_empty, page_count)
            errors.append(err)
            warnings.append(err.cli_message())
        
    except ImportError:
        err = E104_pymupdf_not_installed(str(path))
        errors.append(err)
        warnings.append(err.cli_message())
    except Exception as e:
        err = E106_extractor_exception(str(path), "PyMuPDF", str(e))
        errors.append(err)
        warnings.append(err.cli_message())
    
    # Fallback to pdfplumber if no text extracted
    if not text.strip():
        try:
            import pdfplumber
            
            with pdfplumber.open(path) as pdf:
                if not page_count:
                    page_count = len(pdf.pages)
                pages = []
                for page in pdf.pages:
                    pages.append(page.extract_text() or "")
                    
                    # Extract tables
                    for table in page.extract_tables():
                        if table:
                            md_table = _table_to_markdown(table)
                            tables.append(ExtractedTable(
                                sheet_name=None,
                                data=[],
                                markdown=md_table,
                            ))
                
                text = "\n\n".join(pages)
        
        except ImportError:
            err = E105_pdfplumber_not_installed(str(path))
            errors.append(err)
            warnings.append(err.cli_message())
        except Exception as e:
            err = E106_extractor_exception(str(path), "pdfplumber", str(e))
            errors.append(err)
            warnings.append(err.cli_message())
    
    # OCR fallback if still empty
    if not text.strip():
        try:
            text = extract_pdf_ocr(path)
            method = "ocr"
        except ImportError as e:
            err = E102_ocr_unavailable(str(path), str(e))
            errors.append(err)
            warnings.append(err.cli_message())
        except Exception as e:
            err = E102_ocr_unavailable(str(path), str(e))
            errors.append(err)
            warnings.append(err.cli_message())
    
    # If still no content after all fallbacks, diagnose the root cause
    if not text.strip() and not tables:
        # Check for zlib stream corruption (the specific case from the
        # corrupt Hacker's Guide PDF)
        try:
            import fitz
            doc = fitz.open(path)
            if doc.page_count > 0:
                page = doc[0]
                imgs = page.get_images(full=True)
                if imgs:
                    try:
                        raw = doc.extract_image(imgs[0][0])
                        # If raw image bytes start with damaged data,
                        # this is stream corruption
                        if raw and raw.get("image"):
                            header = raw["image"][:4]
                            # Valid JPEG: FF D8 FF; Valid PNG: 89 50 4E 47
                            if (header[:2] != b'\xff\xd8' and
                                    header[:4] != b'\x89PNG'):
                                zlib_errors_detected = True
                    except Exception:
                        zlib_errors_detected = True
            doc.close()
        except Exception:
            pass
        
        if zlib_errors_detected:
            err = E101_pdf_stream_corruption(str(path), page_count)
        else:
            err = E100_no_content_extracted(str(path))
        errors.append(err)
        warnings.append(err.cli_message())
    
    return ExtractedDoc(
        text=text,
        title=path.stem,
        source_path=path,
        detected_type='pdf',
        extraction_method=method,
        tables=tables,
        warnings=warnings,
        structured_errors=errors,
    )


def extract_pdf_ocr(path: Path) -> str:
    """Extract text from PDF using OCR (Tesseract)."""
    try:
        import pytesseract
        from PIL import Image
        import fitz  # PyMuPDF for page -> image
        
        doc = fitz.open(path)
        pages = []
        
        for page_num in range(len(doc)):
            page = doc[page_num]
            pix = page.get_pixmap()
            img = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
            text = pytesseract.image_to_string(img)
            pages.append(text)
        
        doc.close()
        return "\n\n".join(pages)
    
    except ImportError:
        raise ImportError(
            "OCR dependencies not installed. "
            "Install with: pip install 'lfl[ocr]'"
        )


def extract_docx(path: Path) -> ExtractedDoc:
    """Extract text and tables from DOCX."""
    try:
        from docx import Document
    except ImportError:
        err = E110_docx_not_installed()
        raise ImportError(err.cli_message()) from None
    
    doc = Document(path)
    
    # Extract paragraphs
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    text = "\n\n".join(paragraphs)
    
    # Extract tables
    tables = []
    for table in doc.tables:
        rows = []
        for row in table.rows:
            rows.append([cell.text for cell in row.cells])
        
        if rows:
            md_table = _table_to_markdown(rows)
            tables.append(ExtractedTable(
                sheet_name=None,
                data=[],
                markdown=md_table,
            ))
    
    return ExtractedDoc(
        text=text,
        title=path.stem,
        source_path=path,
        detected_type='docx',
        extraction_method='native',
        tables=tables,
    )


def extract_xlsx(path: Path) -> ExtractedDoc:
    """Extract tables from Excel."""
    try:
        import pandas as pd
    except ImportError:
        err = E111_pandas_not_installed()
        raise ImportError(err.cli_message()) from None
    
    sheets = pd.read_excel(path, sheet_name=None, engine='openpyxl')
    
    tables = []
    for sheet_name, df in sheets.items():
        if not df.empty:
            tables.append(ExtractedTable(
                sheet_name=sheet_name,
                data=df.to_dict('records'),
                markdown=df.to_markdown(index=False),
            ))
    
    # Generate text summary
    text_parts = []
    for table in tables:
        text_parts.append(f"# Sheet: {table.sheet_name}\n\n{table.markdown}")
    
    text = "\n\n".join(text_parts)
    
    return ExtractedDoc(
        text=text,
        title=path.stem,
        source_path=path,
        detected_type='xlsx',
        extraction_method='native',
        tables=tables,
    )


def extract_html(path: Path) -> ExtractedDoc:
    """Extract text from HTML."""
    try:
        from bs4 import BeautifulSoup
    except ImportError:
        err = E112_bs4_not_installed()
        raise ImportError(err.cli_message()) from None
    
    with open(path, 'r', encoding='utf-8') as f:
        html = f.read()
    
    soup = BeautifulSoup(html, 'lxml')
    
    # Remove scripts and styles
    for tag in soup(['script', 'style', 'nav', 'footer']):
        tag.decompose()
    
    # Get title
    title = None
    if soup.title:
        title = soup.title.string
    
    # Extract text
    text = soup.get_text(separator='\n\n')
    
    # Clean up whitespace
    lines = [line.strip() for line in text.split('\n')]
    text = '\n'.join(line for line in lines if line)
    
    return ExtractedDoc(
        text=text,
        title=title or path.stem,
        source_path=path,
        detected_type='html',
        extraction_method='native',
    )


def extract_plain_text(path: Path) -> ExtractedDoc:
    """Extract from plain text files."""
    errors: list[IngestionError] = []
    warnings: list[str] = []
    used_encoding = "utf-8"
    
    # Try different encodings
    for encoding in ['utf-8', 'latin-1', 'cp1252']:
        try:
            with open(path, 'r', encoding=encoding) as f:
                text = f.read()
            used_encoding = encoding
            break
        except UnicodeDecodeError:
            continue
    else:
        # Last resort: binary read with replace
        with open(path, 'rb') as f:
            text = f.read().decode('utf-8', errors='replace')
        used_encoding = "binary-replace"
    
    if used_encoding != "utf-8":
        err = E113_encoding_fallback(str(path), used_encoding)
        errors.append(err)
        warnings.append(err.cli_message())
    
    return ExtractedDoc(
        text=text,
        title=path.stem,
        source_path=path,
        detected_type='text',
        extraction_method='native',
        warnings=warnings,
        structured_errors=errors,
    )


def _table_to_markdown(rows: list[list[str]]) -> str:
    """Convert table rows to Markdown format."""
    if not rows:
        return ""
    
    # Header
    header = "| " + " | ".join(str(cell) for cell in rows[0]) + " |"
    separator = "| " + " | ".join("---" for _ in rows[0]) + " |"
    
    # Body
    body = []
    for row in rows[1:]:
        body.append("| " + " | ".join(str(cell) for cell in row) + " |")
    
    return "\n".join([header, separator] + body)
